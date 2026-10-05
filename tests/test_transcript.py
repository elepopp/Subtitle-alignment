import numpy as np

from subalign.asr.base import Segment, Transcript, Word
from subalign.asr.hallucination import screen
from subalign.formats import write
from subalign.models import Document, Line, Token
from subalign.text.tokenize import make_line


class _Feats:
    """Minimal stand-in for audio Features: voice activity per 10 ms frame."""

    def __init__(self, voiced_spans, total=60.0):
        self.hop_s = 0.01
        self.n = int(total / self.hop_s)
        self.active = np.zeros(self.n)
        for a, b in voiced_spans:
            self.active[int(a / self.hop_s):int(b / self.hop_s)] = 1.0


def _seg(a, b, text, lp=-0.2, nsp=0.1):
    return Segment(a, b, text, [Word(text, a, b, 0.9)], avg_logprob=lp, no_speech_prob=nsp)


def test_hallucination_screening():
    tr = Transcript([
        _seg(0, 4, "大家好，今天聊聊人工智能"),
        _seg(10, 14, "字幕由Amara.org社区提供"),             # credit phrase: removed on its own
        _seg(20, 24, "我们继续今天的话题吧"),                  # text over music (no voice): removed
        _seg(30, 33, "谢谢大家观看"),                          # outro phrase while someone speaks: kept, flagged
        _seg(40, 44, "好的好的好的好的好的好的"),              # looped
    ])
    feats = _Feats([(0, 4), (10, 14), (30, 33), (40, 44)])
    kept, notes = screen(tr, feats)
    texts = [s.text for s in kept.segments]
    assert "大家好，今天聊聊人工智能" in texts and "谢谢大家观看" in texts
    removed = {n["text"] for n in notes if n["action"] == "removed"}
    assert removed == {"字幕由Amara.org社区提供", "我们继续今天的话题吧"}
    flagged = {n["text"] for n in notes if n["action"] == "flagged"}
    assert "谢谢大家观看" in flagged and "好的好的好的好的好的好的" in flagged


def test_songs_may_repeat():
    tr = Transcript([_seg(0, 6, "我爱你我爱你我爱你我爱你")])
    kept, notes = screen(tr, _Feats([(0, 6)]), song=True)
    assert len(kept.segments) == 1 and not notes


def _timed_line(text, t0, step=0.25):
    ln = make_line(text)
    for k, t in enumerate(ln.tokens):
        t.start, t.end = t0 + k * step, t0 + (k + 1) * step
    ln.update_bounds()
    return ln


def test_diarize_splits_lines_at_speaker_changes(monkeypatch):
    import subalign.diarize as D

    # line 2 holds a speaker change after a pause: "好的谢谢" (B) [pause] "那我们开始" (A)
    l1 = _timed_line("欢迎来到节目", 0.0)
    l2 = _timed_line("好的谢谢", 3.0)
    tail = _timed_line("那我们开始", 4.5)
    l2.tokens += tail.tokens
    l2.update_bounds()
    doc = Document(lines=[l1, l2, _timed_line("我也很期待", 8.0)])
    voices = {0.0: [1, 0], 3.0: [0, 1], 4.5: [1, 0], 8.0: [0, 1]}

    def fake_embed(clips, device=None):
        return np.array([voices[k] for k in sorted(voices)][:len(clips)], dtype=float)

    monkeypatch.setattr(D, "embed", fake_embed)
    n = D.diarize(doc, np.zeros(16000 * 12, np.float32), min_chunk=0.5)
    assert n == 2
    assert [(ln.speaker, ln.text) for ln in doc.lines] == [
        ("S1", "欢迎来到节目"), ("S2", "好的谢谢"), ("S1", "那我们开始"), ("S2", "我也很期待")]


def test_transcript_markdown_and_docx():
    a = _timed_line("欢迎收听本期节目", 0.0)
    a.speaker = "S1"
    b = _timed_line("今天聊聊大模型", 2.2)          # same speaker, short pause: same paragraph
    b.speaker = "S1"
    c = _timed_line("大家好我是张伟", 6.0)
    c.speaker = "S2"
    doc = Document(lines=[a, b, c], metadata={"speakers": 2,
                                               "review": [{"start": 6.0, "end": 6.5, "text": "大家", "other": "打架"}]})
    md = write(doc, "transcript")
    assert "**S1**  `00:00`" in md and "**S2**  `00:06`" in md
    assert "欢迎收听本期节目，今天聊聊大模型" in md          # joined with a comma (short pause)
    assert "2 位说话人" in md and "另一模型：「打架」" in md
    data = write(doc, "docx")
    assert isinstance(data, bytes) and data[:2] == b"PK"
    import io

    from docx import Document as Docx

    d = Docx(io.BytesIO(data))
    hl = [r.text for p in d.paragraphs for r in p.runs if r.font.highlight_color]
    assert "".join(hl) == "大家"


def test_cross_check_marks_disagreements(monkeypatch):
    from subalign.align import aligner
    from subalign.align.aligner import AlignConfig
    import subalign.asr as asr_mod

    doc = Document(lines=[_timed_line("今天的天气很好", 0.0)], language="zh")

    class Other:
        def transcribe(self, path, language=None, **kw):
            return Transcript([Segment(0, 2, "今天的天汽很好", [Word("今天的天汽很好", 0, 2)])], "zh")

    monkeypatch.setattr(asr_mod, "get_backend", lambda name, **kw: Other())
    notes = aligner._cross_check(doc, AlignConfig(cross_check="funasr"), "x.wav")
    assert notes["disagreements"] and notes["disagreements"][0]["text"] == "气"
    assert notes["disagreements"][0]["kind"] == "homophone"
    assert doc.lines[0].tokens[4].confidence <= 0.45
