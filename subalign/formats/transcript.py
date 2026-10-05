"""Transcript documents (逐字稿): paragraphs by speaker turn, each with a timestamp.

* ``transcript`` -> Markdown (``.transcript.md``)
* ``docx``       -> Word (``.docx``, needs ``python-docx``); places the two
  recognisers disagreed on (``metadata["review"]``) are highlighted.

A paragraph is a run of lines by the same speaker; it is closed at a speaker
change, at a pause of ``para_gap`` seconds or when it grows past ``max_chars``.
"""
from __future__ import annotations

import io
from typing import Dict, List, Optional, Tuple

from ..models import Document, Token
from ..text.tokenize import is_cjk


def _ts(t: Optional[float]) -> str:
    if t is None:
        return "--:--"
    t = max(0.0, t)
    h, rem = divmod(int(t), 3600)
    m, s = divmod(rem, 60)
    return f"{h:02d}:{m:02d}:{s:02d}" if h else f"{m:02d}:{s:02d}"


def _join(a: str, b: str) -> str:
    if not a or not b:
        return a or b
    return a + b if is_cjk(a[-1]) or is_cjk(b[0]) else a + " " + b


_PUNCT = set("，。！？；：、,.!?;:…—）)」』”")


def paragraphs(doc: Document, para_gap: float = 2.0, max_chars: int = 300) -> List[Dict]:
    out: List[Dict] = []
    for ln in doc.lines:
        if not ln.tokens:
            continue
        cur = out[-1] if out else None
        gap = (ln.start - cur["end"]) if cur and ln.start is not None and cur["end"] is not None else 0.0
        if cur is None or cur["speaker"] != ln.speaker or gap >= para_gap or len(cur["text"]) >= max_chars:
            out.append({"speaker": ln.speaker, "start": ln.start, "end": ln.end, "text": "", "lines": []})
            cur = out[-1]
        text = ln.text.strip()
        prev = cur["text"]
        if prev and text and is_cjk(prev[-1]) and prev[-1] not in _PUNCT:
            # recogniser lines often end without punctuation: mark the boundary by the pause
            prev += "。" if gap >= 0.8 else "，"
        cur["text"] = _join(prev, text)
        cur["lines"].append(ln)
        cur["end"] = ln.end if ln.end is not None else cur["end"]
    return out


def _header(doc: Document) -> List[str]:
    md = doc.metadata or {}
    title = md.get("title") or "逐字稿"
    dur = max((ln.end for ln in doc.lines if ln.end is not None), default=0.0)
    info = [f"时长 {_ts(dur)}"]
    if md.get("speakers"):
        info.append(f"{md['speakers']} 位说话人")
    if md.get("review"):
        info.append(f"{len(md['review'])} 处待核对")
    return [title, " · ".join(info)]


def write_transcript_md(doc: Document, **_) -> str:
    title, info = _header(doc)
    rows = [f"# {title}", "", f"> {info}", ""]
    for p in paragraphs(doc):
        who = f"**{p['speaker']}**  " if p["speaker"] else ""
        rows += [f"{who}`{_ts(p['start'])}`", "", p["text"], ""]
    review = (doc.metadata or {}).get("review") or []
    if review:
        rows += ["---", "", "## 待核对（两个识别模型结果不一致）", ""]
        for r in review:
            other = r.get("other") or "（未识别）"
            rows.append(f"- `{_ts(r['start'])}` 「{r.get('text') or '（无）'}」 / 另一模型：「{other}」")
        rows.append("")
    return "\n".join(rows)


def _review_spans(doc: Document) -> List[Tuple[float, float]]:
    return [(r["start"], r["end"]) for r in (doc.metadata or {}).get("review") or [] if r.get("text")]


def write_transcript_docx(doc: Document, **_) -> bytes:
    try:
        from docx import Document as Docx  # type: ignore
        from docx.enum.text import WD_COLOR_INDEX  # type: ignore
        from docx.oxml.ns import qn  # type: ignore
        from docx.shared import Pt, RGBColor  # type: ignore
    except ImportError as e:  # pragma: no cover
        raise ImportError("docx export needs `pip install python-docx`") from e
    d = Docx()
    st = d.styles["Normal"]
    st.font.name = "Microsoft YaHei"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    st.font.size = Pt(11)
    title, info = _header(doc)
    d.add_heading(title, level=1)
    meta = d.add_paragraph(info)
    meta.runs[0].font.color.rgb = RGBColor(0x66, 0x70, 0x85)
    spans = _review_spans(doc)

    def flagged(t: Token) -> bool:
        return t.start is not None and any(a - 1e-3 <= t.start < b - 1e-3 for a, b in spans)

    for p in paragraphs(doc):
        head = d.add_paragraph()
        if p["speaker"]:
            r = head.add_run(p["speaker"] + "  ")
            r.bold = True
        r = head.add_run(_ts(p["start"]))
        r.font.color.rgb = RGBColor(0x98, 0xA2, 0xB3)
        r.font.size = Pt(9)
        head.paragraph_format.space_after = Pt(0)
        body = d.add_paragraph()
        body.paragraph_format.space_after = Pt(10)
        prev = ""
        for ln in p["lines"]:
            if prev and not (is_cjk(prev[-1:] or "a") or is_cjk((ln.tokens[0].text or "a")[0])):
                body.add_run(" ")
            for i, t in enumerate(ln.tokens):
                run = body.add_run(t.text + (" " if t.space_after and i < len(ln.tokens) - 1 else ""))
                if flagged(t):
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            prev = ln.tokens[-1].text if ln.tokens else prev
    review = (doc.metadata or {}).get("review") or []
    if review:
        d.add_heading("待核对（两个识别模型结果不一致，正文中黄色标出）", level=2)
        for r in review:
            d.add_paragraph(f"{_ts(r['start'])}  「{r.get('text') or '（无）'}」 / 另一模型：「{r.get('other') or '（未识别）'}」",
                            style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
